import hashlib
import io

from langchain_core.documents import Document
from starlette.requests import Request

from app.core.config import settings
from app.main import _client_key
from app.services import document_service as docs
from app.services import ingestion_service


def _ready(ctx, name, content_hash):
    record = docs.create_document(
        ctx, original_name=name, stored_name=name, byte_size=10, content_hash=content_hash
    )
    docs.set_status(record["id"], "ready", pages=1, chunk_count=1)
    return docs.get_document(ctx, record["id"])


def test_duplicate_upload_does_not_overwrite_existing_file(client, user_a):
    ctx = user_a["ctx"]
    original = ctx.storage.upload_path("bob.pdf")
    original.parent.mkdir(parents=True, exist_ok=True)
    original.write_bytes(b"%PDF-1.4 original bob")
    _ready(ctx, "bob.pdf", hashlib.sha256(b"%PDF-1.4 original bob").hexdigest())

    other = b"%PDF-1.4 physics content"
    _ready(ctx, "physics.pdf", hashlib.sha256(other).hexdigest())

    response = client.post(
        "/ingest/",
        headers=user_a["headers"],
        files={"file": ("bob.pdf", io.BytesIO(other), "application/pdf")},
    )
    assert response.status_code == 200
    assert response.json()["status"] == "duplicate"
    assert original.read_bytes() == b"%PDF-1.4 original bob"
    assert not list(original.parent.glob("*.part"))


def test_reingest_removes_previous_vectors(user_a, monkeypatch):
    ctx = user_a["ctx"]
    record = _ready(ctx, "notes.pdf", "hash-notes")
    path = ctx.storage.upload_path("notes.pdf")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"placeholder")
    docs.replace_chunks(record["id"], [{"chunk_id": "old-1", "page": 1}, {"chunk_id": "old-2", "page": 2}])

    deleted = []
    monkeypatch.setattr(
        ingestion_service.faiss_store, "delete_chunks", lambda storage, ids: deleted.extend(ids) or len(ids)
    )
    monkeypatch.setattr(ingestion_service.faiss_store, "add_chunks", lambda storage, chunks: len(chunks))
    monkeypatch.setattr(
        ingestion_service,
        "_load_documents",
        lambda p: [Document(page_content="Fresh content about enzymes.", metadata={"source": "notes.pdf", "page": 1})],
    )

    ingestion_service.ingest_document(ctx, record["id"])

    assert sorted(deleted) == ["old-1", "old-2"]
    new_ids = docs.get_chunk_ids(record["id"])
    assert new_ids and "old-1" not in new_ids


def test_pdf_and_docx_pages_are_one_based(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "notes.docx"
    doc = DocxDocument()
    doc.add_heading("First", level=1)
    doc.add_paragraph("Alpha")
    doc.add_heading("Second", level=1)
    doc.add_paragraph("Beta")
    doc.save(path)

    from app.rag.chunking import chunk_documents

    chunks = chunk_documents(ingestion_service._load_documents(str(path)))
    assert [c.metadata["page"] for c in chunks] == [1, 2]


def test_selection_that_matches_nothing_is_rejected(client, user_a):
    _ready(user_a["ctx"], "notes.pdf", "hash-1")
    response = client.post(
        "/qa/ask",
        headers=user_a["headers"],
        json={"question": "What is photosynthesis?", "document_ids": ["does-not-exist"]},
    )
    assert response.status_code == 400


def test_unknown_session_is_rejected_before_generation(client, user_a, monkeypatch):
    _ready(user_a["ctx"], "notes.pdf", "hash-1")

    def fail(*args, **kwargs):
        raise AssertionError("pipeline should not run")

    monkeypatch.setattr("app.api.routes.qa.run_rag_stages", fail)
    response = client.post(
        "/qa/ask",
        headers=user_a["headers"],
        json={"question": "What is photosynthesis?", "session_id": "missing"},
    )
    assert response.status_code == 404


def test_interrupted_documents_are_marked_failed(user_a):
    record = docs.create_document(
        user_a["ctx"], original_name="a.pdf", stored_name="a.pdf", byte_size=1, content_hash="x"
    )
    docs.set_status(record["id"], "embedding")

    assert docs.fail_interrupted_documents() == 1
    assert docs.get_document(user_a["ctx"], record["id"])["status"] == "failed"


def _request(forwarded):
    return Request(
        {
            "type": "http",
            "headers": [(b"x-forwarded-for", forwarded.encode())],
            "client": ("203.0.113.9", 1234),
        }
    )


def test_forwarded_header_ignored_unless_proxy_trusted(monkeypatch):
    monkeypatch.setattr(settings, "TRUST_PROXY", False)
    assert _client_key(_request("1.1.1.1")) == "203.0.113.9"

    monkeypatch.setattr(settings, "TRUST_PROXY", True)
    assert _client_key(_request("1.1.1.1, 198.51.100.7")) == "198.51.100.7"
