import hashlib
import logging
import os
from pathlib import Path

from app.core.context import UserContext
from app.rag.chunking import chunk_documents
from app.services import document_service as docs
from app.vectorstore import faiss_store

logger = logging.getLogger(__name__)


class IngestionError(RuntimeError):
    pass


def _load_pdf(path: str):
    from langchain_core.documents import Document
    from pypdf import PdfReader

    reader = PdfReader(path)
    source = os.path.basename(path)
    return [
        Document(page_content=page.extract_text() or "", metadata={"source": source, "page": number})
        for number, page in enumerate(reader.pages, start=1)
    ]


def _load_docx(path: str):
    from docx import Document as DocxDocument
    from langchain_core.documents import Document

    doc = DocxDocument(path)
    sections: list[list[str]] = [[]]

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        is_heading = style.startswith("heading") or style == "title"
        if is_heading and sections[-1]:
            sections.append([])
        sections[-1].append(text)

    documents = []
    for index, body in enumerate(sections, start=1):
        if not body:
            continue
        documents.append(
            Document(
                page_content="\n".join(body),
                metadata={"source": os.path.basename(path), "page": index},
            )
        )

    return documents


def _load_documents(persistent_path: str):
    lowered = persistent_path.lower()
    if lowered.endswith(".pdf"):
        return _load_pdf(persistent_path)
    if lowered.endswith(".docx"):
        return _load_docx(persistent_path)
    raise IngestionError("Unsupported file format")


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def ingest_document(ctx: UserContext, document_id: str) -> dict:
    record = docs.get_document(ctx, document_id)
    if record is None:
        raise IngestionError("Document not found")

    storage = ctx.storage
    path = storage.upload_path(record["stored_name"])
    if not path.exists():
        docs.set_status(document_id, "failed", error="Uploaded file is missing")
        raise IngestionError("Uploaded file is missing")

    try:
        stale = docs.get_chunk_ids(document_id)
        if stale:
            faiss_store.delete_chunks(storage, stale)
            docs.replace_chunks(document_id, [])

        docs.set_status(document_id, "parsing")
        loaded = _load_documents(str(path))
        if not loaded:
            raise IngestionError(
                "No text could be extracted. If this is a scanned PDF it needs OCR first."
            )
        pages = len(loaded)

        chunks = chunk_documents(
            loaded,
            document_id=document_id,
            display_name=record["original_name"],
        )
        if not chunks:
            raise IngestionError("Document produced no usable text chunks")

        docs.set_status(document_id, "embedding", pages=pages)
        added = faiss_store.add_chunks(
            storage,
            chunks,
            on_rebuild=lambda: docs.mark_reindex_required(ctx, exclude_id=document_id),
        )

        docs.replace_chunks(
            document_id,
            [
                {
                    "chunk_id": c.metadata["chunk_id"],
                    "page": c.metadata.get("page", 1),
                    "ordinal": c.metadata.get("ordinal", i),
                }
                for i, c in enumerate(chunks)
            ],
        )
        docs.set_status(document_id, "ready", pages=pages, chunk_count=len(chunks))
        logger.info(
            "Indexed document %s for user %s: %s pages, %s chunks (%s new vectors)",
            document_id,
            ctx.user_id,
            pages,
            len(chunks),
            added,
        )
        return {"pages": pages, "chunks": len(chunks), "added": added}

    except IngestionError as exc:
        docs.set_status(document_id, "failed", error=str(exc))
        raise
    except Exception as exc:
        logger.exception("Ingestion failed for document %s", document_id)
        docs.set_status(
            document_id, "failed", error="Processing failed. Please retry this document."
        )
        raise IngestionError("Processing failed. Please retry this document.") from exc


def remove_document(ctx: UserContext, document_id: str) -> bool:
    record = docs.get_document(ctx, document_id)
    if record is None:
        return False

    chunk_ids = docs.get_chunk_ids(document_id)
    if chunk_ids:
        try:
            removed = faiss_store.delete_chunks(ctx.storage, chunk_ids)
            logger.info("Removed %s vectors for document %s", removed, document_id)
        except Exception:
            logger.exception("Vector cleanup failed for document %s", document_id)

    try:
        path = ctx.storage.upload_path(record["stored_name"])
        if path.exists():
            path.unlink()
    except OSError:
        logger.warning("Could not delete file for document %s", document_id)

    return docs.delete_document(ctx, document_id)


def reset_user_documents(ctx: UserContext) -> int:
    removed = docs.delete_all_documents(ctx)
    faiss_store.reset_index(ctx.storage)
    uploads = ctx.storage.uploads_dir
    if uploads.exists():
        for entry in uploads.iterdir():
            if entry.is_file():
                try:
                    entry.unlink()
                except OSError:
                    logger.warning("Could not delete %s", entry.name)
    return removed
